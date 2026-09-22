# -*- coding: utf-8 -*-
"""Spatial-error robustness for the fully adjusted SEE/CIE regressions."""
from __future__ import annotations

import warnings
from pathlib import Path

import numpy as np
import pandas as pd

from config import (
    BASE_YEAR,
    END_YEAR,
    SEE_CIE_REGRESSION_ROOT,
    CITY_SHP,
    RESULT_ROOT,
    CITY_ORDER_4,
)
from utils.extended_analysis import p_to_star
from utils.regression import load_standardized_see_cie_panel

# -----------------------------------------------------------------------------
# Settings
# -----------------------------------------------------------------------------
K_NEIGHBORS = 4
N_PERMUTATIONS = 9999
RANDOM_SEED = 20260828
LAMBDA_BOUND = 0.99
HESSIAN_EPS = 1e-4

OUT_ROOT = RESULT_ROOT / "spatial_robustness"

ACC_BASE_COL = f"acc_{BASE_YEAR}"
GINI_BASE_COL = f"gini_{BASE_YEAR}"
THEIL_BASE_COL = f"theil_{BASE_YEAR}"
ATKINSON_BASE_COL = f"atkinson_05_{BASE_YEAR}"
GDP_BASE_COL = f"GDP_{BASE_YEAR}"
RESPOP_BASE_COL = f"ResPop_{BASE_YEAR}"
FISCAL_BASE_COL = f"ln_FiscalRevenue_pc_{BASE_YEAR}"


def _norm_text(s: pd.Series) -> pd.Series:
    return s.astype("string").str.strip()


def build_city_geometry():
    """Load the fixed city boundary layer used for spatial robustness analyses."""
    try:
        import geopandas as gpd
    except ImportError as exc:
        raise ImportError("geopandas is required for spatial robustness analyses") from exc

    if not Path(CITY_SHP).exists():
        raise FileNotFoundError(f"CITY_SHP not found: {CITY_SHP}")
    gdf = gpd.read_file(CITY_SHP)
    required = {"地名", "地级", "geometry"}
    missing = required - set(gdf.columns)
    if missing:
        raise ValueError(f"CITY_SHP is missing required columns: {sorted(missing)}")

    gdf = gdf[["地名", "地级", "geometry"]].copy()
    gdf["city_name"] = _norm_text(gdf["地级"])
    fallback = gdf["city_name"].eq("不统计") | gdf["city_name"].isna()
    gdf.loc[fallback, "city_name"] = _norm_text(gdf.loc[fallback, "地名"])

    bad = ~gdf.geometry.is_valid
    if bad.any():
        warnings.warn(f"Repairing {int(bad.sum())} invalid city geometries with buffer(0).")
        gdf.loc[bad, "geometry"] = gdf.loc[bad, "geometry"].buffer(0)
    return gdf[["city_name", "geometry"]].drop_duplicates("city_name")


def prepare_regression_data() -> pd.DataFrame:
    """Load the standardized city panel shared with the main regressions."""
    df = load_standardized_see_cie_panel()
    needed = {"地级", "city_level", FISCAL_BASE_COL}
    missing = needed - set(df.columns)
    if missing:
        raise ValueError(f"Regression dataset missing columns: {sorted(missing)}")
    df["地级"] = _norm_text(df["地级"])
    return df

def stage3_specs():
    """Return the seven fiscal-adjusted Stage-3 formulas used in main regression."""
    city_fe = 'C(city_level_4, Treatment(reference="Medium/Small City"))'
    controls = (
        f"pop_density_mean + {GDP_BASE_COL} + GDP_growth_pct + {RESPOP_BASE_COL} "
        "+ Ppo_NetIn_rate + ResPop_growth_rate + " + FISCAL_BASE_COL
    )

    specs = [
        {
            "model": "accessibility_SEE_CIE_stage3",
            "outcome_label": "accessibility",
            "outcome": "acc_delta",
            "family": "SEE_CIE",
            "formula": (
                f"acc_delta ~ {ACC_BASE_COL} + {controls} "
                f"+ city_SEE*{city_fe} + city_CIE*{city_fe}"
            ),
            "columns": [
                "acc_delta", ACC_BASE_COL, "pop_density_mean", GDP_BASE_COL,
                "GDP_growth_pct", RESPOP_BASE_COL, "Ppo_NetIn_rate",
                "ResPop_growth_rate", FISCAL_BASE_COL, "city_SEE", "city_CIE",
                "city_level_4",
            ],
            "expansions": [("SEE", "city_SEE"), ("CIE", "city_CIE")],
        }
    ]

    inequality = {
        "gini": ("gini_delta", GINI_BASE_COL),
        "theil": ("theil_delta", THEIL_BASE_COL),
        "atkinson_05": ("atkinson_05_delta", ATKINSON_BASE_COL),
    }
    for label, (y, y0) in inequality.items():
        common_cols = [
            y, y0, "pop_density_mean", GDP_BASE_COL, "GDP_growth_pct",
            RESPOP_BASE_COL, "Ppo_NetIn_rate", "ResPop_growth_rate",
            FISCAL_BASE_COL, "city_level_4",
        ]
        specs.append(
            {
                "model": f"{label}_SEE_CIE_stage3",
                "outcome_label": label,
                "outcome": y,
                "family": "SEE_CIE",
                "formula": (
                    f"{y} ~ {y0} + {controls} "
                    f"+ city_SEE*{city_fe} + city_CIE*{city_fe}"
                ),
                "columns": common_cols + ["city_SEE", "city_CIE"],
                "expansions": [("SEE", "city_SEE"), ("CIE", "city_CIE")],
            }
        )
        specs.append(
            {
                "model": f"{label}_TotalExpansion_stage3",
                "outcome_label": label,
                "outcome": y,
                "family": "TotalExpansion",
                "formula": (
                    f"{y} ~ {y0} + {controls} "
                    f"+ city_TotalExpansion*{city_fe}"
                ),
                "columns": common_cols + ["city_TotalExpansion"],
                "expansions": [("TotalExpansion", "city_TotalExpansion")],
            }
        )
    return specs


def build_knn_weights(city_geo, city_order, k=4):
    """Build row-standardized KNN weights for the spatial analysis."""
    try:
        from libpysal.weights import KNN
    except ImportError as e:
        raise ImportError("Please install libpysal: pip install libpysal") from e

    q = city_geo.set_index("city_name").loc[list(city_order)].copy()
    if q.crs is None:
        raise ValueError("CITY_SHP has no CRS; cannot construct defensible KNN distances.")

    china_aea = (
        "+proj=aea +lat_1=25 +lat_2=47 +lat_0=0 +lon_0=105 "
        "+datum=WGS84 +units=m +no_defs"
    )
    q_proj = q.to_crs(china_aea)
    cent = q_proj.geometry.centroid
    coords = np.column_stack([cent.x.to_numpy(), cent.y.to_numpy()])
    if len(coords) <= k:
        raise ValueError(f"Model sample N={len(coords)} must be greater than k={k}.")

    w = KNN.from_array(coords, k=k, ids=list(city_order))
    w.transform = "R"
    if list(w.id_order) != list(city_order):
        raise RuntimeError("libpysal weight id_order does not match regression row order.")
    return w


def moran_stats(values, w):
    try:
        from esda.moran import Moran
    except ImportError as e:
        raise ImportError("Please install esda: pip install esda") from e

    arr = np.asarray(values, dtype=float).reshape(-1)
    if not np.isfinite(arr).all():
        raise ValueError("Moran input contains NaN or infinite values.")
    if np.nanstd(arr) == 0:
        return {
            "moran_I": np.nan, "expected_I": np.nan, "p_perm": np.nan,
            "z_sim": np.nan, "p_z_sim": np.nan, "permutations": N_PERMUTATIONS,
        }
    np.random.seed(RANDOM_SEED)
    mi = Moran(arr, w, permutations=N_PERMUTATIONS, two_tailed=True)
    return {
        "moran_I": float(mi.I),
        "expected_I": float(mi.EI),
        "p_perm": float(mi.p_sim),
        "z_sim": float(mi.z_sim),
        "p_z_sim": float(mi.p_z_sim),
        "permutations": int(N_PERMUTATIONS),
    }


def _w_dense(w) -> np.ndarray:
    mat, ids = w.full()
    if list(ids) != list(w.id_order):
        raise RuntimeError("Unexpected libpysal full() id ordering.")
    W = np.asarray(mat, dtype=float)
    if not np.isfinite(W).all():
        raise ValueError("Spatial weights contain non-finite values.")
    return W


def fit_sem_ml(y, X, W, coef_names):
    """Gaussian ML spatial-error model with observed-Hessian covariance.

    Returns a dict containing beta, beta covariance, lambda, lambda SE, sigma2,
    log-likelihood and covariance estimates.
    """
    try:
        from scipy.optimize import minimize_scalar
        from scipy.stats import norm
        from statsmodels.tools.numdiff import approx_hess
    except ImportError as e:
        raise ImportError("scipy and statsmodels are required for spatial robustness") from e

    y = np.asarray(y, dtype=float).reshape(-1)
    X = np.asarray(X, dtype=float)
    W = np.asarray(W, dtype=float)
    n, p = X.shape
    if y.shape[0] != n or W.shape != (n, n):
        raise ValueError("SEM dimensions are inconsistent.")
    if len(coef_names) != p:
        raise ValueError("Coefficient-name length does not match X columns.")
    if np.linalg.matrix_rank(X) < p:
        raise np.linalg.LinAlgError(
            f"SEM design matrix is rank deficient: rank={np.linalg.matrix_rank(X)}, p={p}"
        )

    I = np.eye(n, dtype=float)
    eigvals = np.linalg.eigvals(W)

    def _logabsdet(lam: float) -> float:
        # log |det(I-lambda W)| from precomputed eigenvalues.
        vals = 1.0 - float(lam) * eigvals
        absvals = np.abs(vals)
        if np.any(absvals <= 1e-14) or not np.isfinite(absvals).all():
            return -np.inf
        return float(np.sum(np.log(absvals)).real)

    def _concentrated_nll(lam: float, return_parts: bool = False):
        logdet = _logabsdet(lam)
        if not np.isfinite(logdet):
            return (1e100, None, None, None) if return_parts else 1e100
        A = I - float(lam) * W
        yt = A @ y
        Xt = A @ X
        beta = np.linalg.lstsq(Xt, yt, rcond=None)[0]
        eps = yt - Xt @ beta
        sse = float(eps @ eps)
        sigma2 = max(sse / n, 1e-12)
        nll = 0.5 * n * (np.log(2.0 * np.pi) + 1.0 + np.log(sigma2)) - logdet
        if return_parts:
            return float(nll), beta, float(sigma2), eps
        return float(nll)

    opt = minimize_scalar(
        _concentrated_nll,
        bounds=(-LAMBDA_BOUND, LAMBDA_BOUND),
        method="bounded",
        options={"xatol": 1e-8, "maxiter": 500},
    )
    if not opt.success:
        raise RuntimeError(f"SEM lambda optimization failed: {opt.message}")

    lam = float(opt.x)
    nll, beta, sigma2, innovation = _concentrated_nll(lam, return_parts=True)
    if beta is None:
        raise RuntimeError("SEM failed to recover concentrated beta at optimum.")

    # Full likelihood Hessian.  Parameterization psi -> lambda keeps the optimizer/Hessian
    # away from an invalid spatial-autoregressive boundary.
    lam_scaled = np.clip(lam / LAMBDA_BOUND, -0.999999, 0.999999)
    psi = float(np.arctanh(lam_scaled))
    theta = np.concatenate([beta, [np.log(sigma2), psi]])

    def _full_nll(theta_vec):
        b = np.asarray(theta_vec[:p], dtype=float)
        s2 = float(np.exp(theta_vec[p]))
        la = float(LAMBDA_BOUND * np.tanh(theta_vec[p + 1]))
        logdet = _logabsdet(la)
        if not np.isfinite(logdet) or not np.isfinite(s2) or s2 <= 0:
            return 1e100
        A = I - la * W
        e = A @ (y - X @ b)
        return float(
            0.5 * n * (np.log(2.0 * np.pi) + np.log(s2))
            + 0.5 * float(e @ e) / s2
            - logdet
        )

    H = np.asarray(approx_hess(theta, _full_nll, epsilon=HESSIAN_EPS), dtype=float)
    H = 0.5 * (H + H.T)
    cov_theta = np.linalg.pinv(H, rcond=1e-12)
    cov_theta = 0.5 * (cov_theta + cov_theta.T)
    cov_beta = cov_theta[:p, :p]

    beta_var = np.diag(cov_beta)
    if np.any(beta_var < -1e-8):
        warnings.warn(
            "SEM covariance matrix has materially negative coefficient variances."
        )
    beta_se = np.sqrt(np.maximum(beta_var, 0.0))
    z_beta = np.divide(beta, beta_se, out=np.full_like(beta, np.nan), where=beta_se > 0)
    p_beta = 2.0 * norm.sf(np.abs(z_beta))

    dlam_dpsi = LAMBDA_BOUND * (1.0 - np.tanh(psi) ** 2)
    var_psi = float(cov_theta[p + 1, p + 1])
    lambda_se = abs(dlam_dpsi) * np.sqrt(max(var_psi, 0.0))
    lambda_z = lam / lambda_se if lambda_se > 0 else np.nan
    lambda_p = float(2.0 * norm.sf(abs(lambda_z))) if np.isfinite(lambda_z) else np.nan

    residual_u = y - X @ beta
    innovation = (I - lam * W) @ residual_u
    llf = -float(nll)

    return {
        "coef_names": list(coef_names),
        "beta": np.asarray(beta, dtype=float),
        "beta_se": np.asarray(beta_se, dtype=float),
        "beta_p": np.asarray(p_beta, dtype=float),
        "cov_beta": np.asarray(cov_beta, dtype=float),
        "lambda": lam,
        "lambda_se": float(lambda_se),
        "lambda_p": float(lambda_p),
        "sigma2": float(sigma2),
        "llf": llf,
        "residual_u": np.asarray(residual_u, dtype=float),
        "innovation": np.asarray(innovation, dtype=float),
    }


def _find_interaction_name(names, var_prefix, city):
    hits = [
        name for name in names
        if var_prefix in name
        and f"[T.{city}]" in name
        and "C(city_level_4" in name
        and ":" in name
    ]
    return hits[0] if hits else None


def _linear_effect(coef_names, beta, cov, var_prefix, city):
    """City-size-specific main+interaction effect and normal-theory Wald p-value."""
    from scipy.stats import norm

    names = list(coef_names)
    if var_prefix not in names:
        return np.nan, np.nan, np.nan
    r = np.zeros(len(names), dtype=float)
    r[names.index(var_prefix)] = 1.0
    if city != "Medium/Small City":
        inter = _find_interaction_name(names, var_prefix, city)
        if inter is None:
            raise KeyError(f"Missing interaction for {var_prefix} × {city}")
        r[names.index(inter)] = 1.0
    effect = float(r @ beta)
    var = float(r @ cov @ r)
    se = float(np.sqrt(max(var, 0.0)))
    z = effect / se if se > 0 else np.nan
    p = float(2.0 * norm.sf(abs(z))) if np.isfinite(z) else np.nan
    return effect, se, p


def _ols_effect(model, var_prefix, city):
    names = list(model.params.index)
    beta = model.params.to_numpy(dtype=float)
    cov = model.cov_params().to_numpy(dtype=float)
    return _linear_effect(names, beta, cov, var_prefix, city)


def _coef_rows_sem(fit, *, model_name, outcome, family, n):
    rows = []
    for name, b, se, p in zip(
        fit["coef_names"], fit["beta"], fit["beta_se"], fit["beta_p"]
    ):
        rows.append({
            "model": model_name,
            "outcome": outcome,
            "family": family,
            "term": name,
            "coef": float(b),
            "se": float(se),
            "p": float(p),
            "significance": p_to_star(p),
            "n": int(n),
            "estimator": "SEM_ML",
            "weights": f"KNN{K_NEIGHBORS}",
        })
    rows.append({
        "model": model_name,
        "outcome": outcome,
        "family": family,
        "term": "lambda",
        "coef": fit["lambda"],
        "se": fit["lambda_se"],
        "p": fit["lambda_p"],
        "significance": p_to_star(fit["lambda_p"]),
        "n": int(n),
        "estimator": "SEM_ML",
        "weights": f"KNN{K_NEIGHBORS}",
    })
    return rows


def _si_p_to_star(pval: float) -> str:
    """Significance stars used in the supplementary comparison table."""
    try:
        p = float(pval)
    except (TypeError, ValueError):
        return ""
    if not np.isfinite(p):
        return ""
    if p < 0.01:
        return "***"
    if p < 0.05:
        return "**"
    if p < 0.10:
        return "*"
    return ""


def write_supplementary_compact_docx(effects: pd.DataFrame, out_path: Path) -> None:
    """Write a compact Word table comparing OLS and SEM estimates by city size."""
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
    except ImportError as e:
        raise ImportError(
            "python-docx is required to write the Supplementary Word table. "
            "Install it with: pip install python-docx"
        ) from e

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    outcome_labels = {
        "accessibility": "Accessibility improvement",
        "gini": "Gini inequality change",
        "theil": "Theil T inequality change",
        "atkinson_05": "Atkinson (epsilon = 0.5) inequality change",
    }
    expansion_labels = {
        "SEE": "SEE",
        "CIE": "CIE",
        "TotalExpansion": "Total expansion",
    }
    city_order = list(CITY_ORDER_4)
    row_order = [
        ("accessibility", "SEE"),
        ("accessibility", "CIE"),
        ("gini", "SEE"),
        ("gini", "CIE"),
        ("gini", "TotalExpansion"),
        ("theil", "SEE"),
        ("theil", "CIE"),
        ("theil", "TotalExpansion"),
        ("atkinson_05", "SEE"),
        ("atkinson_05", "CIE"),
        ("atkinson_05", "TotalExpansion"),
    ]

    df = effects.copy()
    required = {
        "outcome", "expansion", "city_level", "n",
        "OLS_HC1_effect", "OLS_HC1_se", "OLS_HC1_p",
        "SEM_effect", "SEM_se", "SEM_p",
    }
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Cannot build Supplementary Word table; missing columns: {sorted(missing)}")

    # All rows should use the same analytical sample in the formal Stage-3 models.
    n_values = sorted(pd.to_numeric(df["n"], errors="coerce").dropna().astype(int).unique())
    n_note = str(n_values[0]) if len(n_values) == 1 else ", ".join(map(str, n_values))

    def _set_cell_shading(cell, fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), fill)

    def _set_cell_margins(cell, top=45, start=55, bottom=45, end=55) -> None:
        tc = cell._tc
        tc_pr = tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
            node = tc_mar.find(qn(f"w:{m}"))
            if node is None:
                node = OxmlElement(f"w:{m}")
                tc_mar.append(node)
            node.set(qn("w:w"), str(v))
            node.set(qn("w:type"), "dxa")

    def _set_repeat_table_header(row) -> None:
        tr_pr = row._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)

    def _add_estimate_line(cell, label: str, beta: float, se: float, pval: float) -> None:
        p = cell.add_paragraph() if cell.text else cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(f"{label}: {beta:.3f}")
        r.font.size = Pt(8.2)
        stars = _si_p_to_star(float(pval))
        if stars:
            rs = p.add_run(stars)
            rs.font.size = Pt(7.2)
            rs.font.superscript = True
        r2 = p.add_run(f" ({se:.3f})")
        r2.font.size = Pt(8.2)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.keep_with_next = True
    tr = title.add_run(
        "Comparison of OLS and spatial-error model estimates of "
        "hospital-expansion associations across city-size groups."
    )
    tr.bold = True
    tr.font.name = "Arial"
    tr.font.size = Pt(9.5)

    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    widths = [Inches(2.15), Inches(1.15), Inches(1.60), Inches(1.45), Inches(1.45), Inches(1.45)]
    headers = [
        "Outcome", "Expansion pathway", "Medium/Small City",
        "Large City", "Super City", "Mega City",
    ]
    hdr = table.rows[0]
    _set_repeat_table_header(hdr)
    for j, (cell, text, width) in enumerate(zip(hdr.cells, headers, widths)):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cell, "D9E2F3")
        _set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(8.3)

    row_cells = []
    for outcome, expansion in row_order:
        subset = df[(df["outcome"] == outcome) & (df["expansion"] == expansion)].copy()
        if subset.empty:
            raise ValueError(f"Missing result row for outcome={outcome}, expansion={expansion}")

        row = table.add_row()
        cells = row.cells
        row_cells.append((outcome, cells))
        for j, width in enumerate(widths):
            cells[j].width = width
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cells[j])

        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(outcome_labels[outcome])
        r0.font.name = "Arial"
        r0.font.size = Pt(8.2)

        p1 = cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(expansion_labels[expansion])
        r1.font.name = "Arial"
        r1.font.size = Pt(8.2)

        for k, city in enumerate(city_order, start=2):
            one = subset[subset["city_level"] == city]
            if len(one) != 1:
                raise ValueError(
                    f"Expected one result for outcome={outcome}, expansion={expansion}, city={city}; "
                    f"found {len(one)}"
                )
            x = one.iloc[0]
            cell = cells[k]
            cell.text = ""
            _add_estimate_line(
                cell, "OLS", float(x["OLS_HC1_effect"]),
                float(x["OLS_HC1_se"]), float(x["OLS_HC1_p"])
            )
            _add_estimate_line(
                cell, "SEM", float(x["SEM_effect"]),
                float(x["SEM_se"]), float(x["SEM_p"])
            )

    # Merge the outcome cells to avoid repeating the same label.
    groups = [
        (0, 1),   # accessibility
        (2, 4),   # gini
        (5, 7),   # theil
        (8, 10),  # atkinson
    ]
    for start, end in groups:
        first_cell = table.rows[start + 1].cells[0]
        last_cell = table.rows[end + 1].cells[0]
        merged = first_cell.merge(last_cell)
        merged.text = ""
        merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(merged)
        p = merged.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        outcome_key = row_order[start][0]
        run = p.add_run(outcome_labels[outcome_key])
        run.font.name = "Arial"
        run.font.size = Pt(8.2)

    # Prevent rows from splitting across pages.
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(0)
    note.paragraph_format.keep_together = True
    nr = note.add_run("Note. " )
    nr.bold = True
    nr.font.name = "Arial"
    nr.font.size = Pt(8.2)
    note_text = (
        "Entries are beta coefficients with standard errors in parentheses. Within each city-size cell, "
        "the first line reports heteroskedasticity-robust OLS estimates (HC1) and the second line reports "
        f"maximum-likelihood spatial error model (SEM) estimates using row-standardized KNN-{K_NEIGHBORS} "
        "spatial weights based on prefecture-level city centroids. All models use the same fiscal-adjusted "
        f"Stage-3 analytical sample and mean specification (N = {n_note} cities). City-size-specific effects "
        "are linear combinations of the corresponding main and interaction terms. "
        "* P < 0.10; ** P < 0.05; *** P < 0.01."
    )
    rr = note.add_run(note_text)
    rr.font.name = "Arial"
    rr.font.size = Pt(8.2)

    doc.save(out_path)

def main():
    try:
        import patsy
        import statsmodels.formula.api as smf
        from scipy.stats import chi2
    except ImportError as e:
        raise ImportError("patsy, statsmodels and scipy are required for spatial robustness") from e

    OUT_ROOT.mkdir(parents=True, exist_ok=True)

    df = prepare_regression_data()
    city_geo = build_city_geometry()
    geo_names = set(city_geo["city_name"].astype(str))

    summary_rows = []
    coef_rows = []
    effect_rows = []
    moran_rows = []

    for spec in stage3_specs():
        required = list(dict.fromkeys(spec["columns"] + ["地级"]))
        missing_cols = [c for c in required if c not in df.columns]
        if missing_cols:
            raise ValueError(f"{spec['model']} missing columns: {missing_cols}")

        d = df[required].dropna().copy()
        d["地级"] = _norm_text(d["地级"])
        dup = d["地级"].duplicated(keep=False)
        if dup.any():
            names = sorted(d.loc[dup, "地级"].astype(str).unique())
            raise ValueError(f"{spec['model']} has duplicate city observations: {names[:20]}")

        matched = d["地级"].isin(geo_names)
        if not matched.all():
            bad = d.loc[~matched, "地级"].astype(str).tolist()
            raise ValueError(
                f"{spec['model']} has {len(bad)} cities unmatched to geometry: {bad}"
            )

        # Build the exact patsy design for the Stage-3 formula.  This also gives us
        # coefficient names identical to statsmodels main regression.
        y_df, X_df = patsy.dmatrices(spec["formula"], data=d, return_type="dataframe")
        used_idx = y_df.index
        d_used = d.loc[used_idx].copy()
        if len(d_used) != len(y_df):
            raise RuntimeError(f"Patsy row alignment failed for {spec['model']}.")

        city_order = d_used["地级"].astype(str).tolist()
        w = build_knn_weights(city_geo, city_order, k=K_NEIGHBORS)
        W = _w_dense(w)

        # Same-sample HC1 OLS comparison.
        ols = smf.ols(spec["formula"], data=d_used).fit(cov_type="HC1")
        if list(ols.model.data.row_labels) != list(d_used.index):
            raise RuntimeError(f"OLS row alignment failed for {spec['model']}.")

        fit = fit_sem_ml(
            y_df.iloc[:, 0].to_numpy(dtype=float),
            X_df.to_numpy(dtype=float),
            W,
            list(X_df.columns),
        )
        n = len(d_used)

        # Gaussian OLS llf versus SEM llf gives a conventional 1-df LR test for lambda.
        lr = max(0.0, 2.0 * (fit["llf"] - float(ols.llf)))
        lr_p = float(chi2.sf(lr, df=1))

        ols_moran = moran_stats(ols.resid.to_numpy(dtype=float), w)
        sem_u_moran = moran_stats(fit["residual_u"], w)
        sem_eps_moran = moran_stats(fit["innovation"], w)

        summary_rows.append({
            "model": spec["model"],
            "outcome": spec["outcome_label"],
            "family": spec["family"],
            "n": n,
            "weights": f"KNN{K_NEIGHBORS}",
            "lambda": fit["lambda"],
            "lambda_se": fit["lambda_se"],
            "lambda_p": fit["lambda_p"],
            "sigma2": fit["sigma2"],
            "llf_OLS": float(ols.llf),
            "llf_SEM": fit["llf"],
            "LR_lambda": lr,
            "LR_lambda_p": lr_p,
            "OLS_resid_moran_I": ols_moran["moran_I"],
            "OLS_resid_moran_p": ols_moran["p_perm"],
            "SEM_innovation_moran_I": sem_eps_moran["moran_I"],
            "SEM_innovation_moran_p": sem_eps_moran["p_perm"],
        })

        coef_rows.extend(_coef_rows_sem(
            fit,
            model_name=spec["model"],
            outcome=spec["outcome_label"],
            family=spec["family"],
            n=n,
        ))

        for expansion, var_prefix in spec["expansions"]:
            for city in CITY_ORDER_4:
                ols_eff, ols_se, ols_p = _ols_effect(ols, var_prefix, city)
                sem_eff, sem_se, sem_p = _linear_effect(
                    fit["coef_names"], fit["beta"], fit["cov_beta"], var_prefix, city
                )
                effect_rows.append({
                    "model": spec["model"],
                    "outcome": spec["outcome_label"],
                    "family": spec["family"],
                    "expansion": expansion,
                    "city_level": city,
                    "n": n,
                    "OLS_HC1_effect": ols_eff,
                    "OLS_HC1_se": ols_se,
                    "OLS_HC1_p": ols_p,
                    "SEM_effect": sem_eff,
                    "SEM_se": sem_se,
                    "SEM_p": sem_p,
                })

        for statistic, vals, mr in [
            ("OLS_residual", ols.resid.to_numpy(dtype=float), ols_moran),
            ("SEM_raw_residual_u", fit["residual_u"], sem_u_moran),
            ("SEM_innovation_epsilon", fit["innovation"], sem_eps_moran),
        ]:
            moran_rows.append({
                "model": spec["model"],
                "outcome": spec["outcome_label"],
                "family": spec["family"],
                "statistic": statistic,
                "n": n,
                "weights": f"KNN{K_NEIGHBORS}",
                "k": K_NEIGHBORS,
                **mr,
            })


    summary = pd.DataFrame(summary_rows)
    coefficients = pd.DataFrame(coef_rows)
    effects = pd.DataFrame(effect_rows)
    moran = pd.DataFrame(moran_rows)
    summary.to_csv(OUT_ROOT / "sem_model_summary.csv", index=False, encoding="utf-8-sig")
    coefficients.to_csv(OUT_ROOT / "sem_coefficients_long.csv", index=False, encoding="utf-8-sig")
    effects.to_csv(OUT_ROOT / "sem_city_specific_effects_long.csv", index=False, encoding="utf-8-sig")
    moran.to_csv(OUT_ROOT / "sem_residual_moran_knn4.csv", index=False, encoding="utf-8-sig")
    # Formatted OLS-versus-SEM coefficient table.
    # Keep beta, SE and exact P values in separate columns.
    outcome_labels = {
        "accessibility": "Accessibility improvement",
        "gini": "Gini inequality change",
        "theil": "Theil T inequality change",
        "atkinson_05": "Atkinson (epsilon=0.5) inequality change",
    }
    expansion_labels = {
        "SEE": "SEE",
        "CIE": "CIE",
        "TotalExpansion": "Total expansion",
    }

    si = effects.copy()
    si["Outcome"] = si["outcome"].map(outcome_labels).fillna(si["outcome"])
    si["Expansion pathway"] = si["expansion"].map(expansion_labels).fillna(si["expansion"])
    si["City-size group"] = si["city_level"]
    si["N"] = si["n"].astype(int)
    si["OLS beta"] = si["OLS_HC1_effect"]
    si["OLS SE"] = si["OLS_HC1_se"]
    si["OLS P"] = si["OLS_HC1_p"]
    si["SEM beta"] = si["SEM_effect"]
    si["SEM SE"] = si["SEM_se"]
    si["SEM P"] = si["SEM_p"]

    # Preserve the manuscript ordering of outcomes, expansion pathways and city sizes.
    outcome_order = [
        "Accessibility improvement",
        "Gini inequality change",
        "Theil T inequality change",
        "Atkinson (epsilon=0.5) inequality change",
    ]
    expansion_order = ["SEE", "CIE", "Total expansion"]
    si["_outcome_order"] = pd.Categorical(si["Outcome"], categories=outcome_order, ordered=True)
    si["_expansion_order"] = pd.Categorical(
        si["Expansion pathway"], categories=expansion_order, ordered=True
    )
    si["_city_order"] = pd.Categorical(
        si["City-size group"], categories=CITY_ORDER_4, ordered=True
    )
    si = si.sort_values(["_outcome_order", "_expansion_order", "_city_order"]).copy()

    si = si[[
        "Outcome", "Expansion pathway", "City-size group", "N",
        "OLS beta", "OLS SE", "OLS P",
        "SEM beta", "SEM SE", "SEM P",
    ]]
    si.to_csv(
        OUT_ROOT / "Supplementary_Table_OLS_vs_SEM_city_effects.csv",
        index=False,
        encoding="utf-8-sig",
        float_format="%.6f",
    )

    # Compact Word version of the same comparison.
    write_supplementary_compact_docx(
        effects, OUT_ROOT / "Supplementary_Table_OLS_vs_SEM_compact.docx"
    )

    readme = f"""Spatial-error robustness for fiscal-adjusted Stage-3 models

Model:
  Gaussian maximum-likelihood spatial error model with row-standardized KNN{K_NEIGHBORS} weights.
  Each SEM uses the same city sample and mean specification as the corresponding Stage-3 HC1 OLS model.

Key files:
  sem_model_summary.csv
  sem_coefficients_long.csv
  sem_city_specific_effects_long.csv
  Supplementary_Table_OLS_vs_SEM_city_effects.csv
  Supplementary_Table_OLS_vs_SEM_compact.docx
  sem_residual_moran_knn4.csv
"""
    (OUT_ROOT / "README.txt").write_text(readme, encoding="utf-8")



if __name__ == "__main__":
    main()
